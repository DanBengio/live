# Résumé Excel des conférences Live : py -3.14 scripts/synthese-excel.py conf-data.json sortie.xlsx
# conf-data.json : liste de { who, code, label, title, createdAt, people, items:[{...activité, answers}] }
# (voir scripts/conferences-data.cjs). Nécessite openpyxl (py -3.14 -m pip install --user openpyxl).
import json
import sys
import unicodedata
from collections import Counter
from datetime import datetime, timezone, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
confs = json.load(open(SRC, encoding="utf-8"))
CASA = timezone(timedelta(hours=1))  # heure du Maroc

FONT = "Arial"
F = lambda **k: Font(name=FONT, **k)
HEAD_FILL = PatternFill("solid", fgColor="1E1B4B")
SUB_FILL = PatternFill("solid", fgColor="E0E7FF")
TOTAL_FILL = PatternFill("solid", fgColor="F1F5F9")
thin = Side(style="thin", color="CBD5E1")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)


def head(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def cell(ws, row, col, value, fmt=None, bold=False, fill=None, wrap=False, align=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = F(bold=bold)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = Alignment(wrap_text=wrap, vertical="center", horizontal=align)
    return c


def norm(name):
    s = unicodedata.normalize("NFD", str(name).strip().lower())
    return " ".join("".join(ch for ch in s if unicodedata.category(ch) != "Mn").split())


def when(ms):
    return datetime.fromtimestamp(ms / 1000, CASA).replace(tzinfo=None)


def word(w):
    return " ".join(str(w).strip().lower().split())


items = confs[0]["items"]  # mêmes questions dans les trois conférences (mêmes identifiants)
for c in confs:
    assert [i["id"] for i in c["items"]] == [i["id"] for i in items], "questions différentes"
N = len(confs)
names = [c["who"] for c in confs]

wb = Workbook()

# ---------------------------------------------------------------- Résultats
res = wb.active
res.title = "Résultats"
res["A1"] = "Résultats par question — toutes conférences"
res["A1"].font = F(bold=True, size=14)
res["A2"] = "Nombre de choix par conférence. % = part des répondants (une personne peut cocher plusieurs choix sur les questions à choix multiples)."
res["A2"].font = F(italic=True, color="64748B")
cols = ["Question / choix"] + names + ["Total", "% des répondants"]
head(res, 4, cols, [62] + [12] * N + [12, 16])
res.freeze_panes = "B5"
r = 5
energy_rows = {}  # clé de question -> (première ligne, dernière ligne) des choix 1..10
answered_cells = {}  # index question -> ligne « répondants »
for qi, it in enumerate(items):
    t = it["type"]
    label = f"Q{qi + 1} · {it['title']}" + ("  (plusieurs choix possibles)" if it.get("multiple") else "")
    cell(res, r, 1, label, bold=True, fill=SUB_FILL, wrap=True)
    for k in range(2, len(cols) + 1):
        cell(res, r, k, None, fill=SUB_FILL)
    r += 1
    # répondants
    resp_row = r
    cell(res, r, 1, "Répondants", bold=True, fill=TOTAL_FILL)
    for j, c in enumerate(confs):
        cell(res, r, 2 + j, len(c["items"][qi]["answers"]), fmt="0", fill=TOTAL_FILL, align="center")
    cell(res, r, 2 + N, f"=SUM({get_column_letter(2)}{r}:{get_column_letter(1 + N)}{r})", fmt="0", bold=True, fill=TOTAL_FILL, align="center")
    cell(res, r, 3 + N, None, fill=TOTAL_FILL)
    answered_cells[qi] = r
    r += 1
    if t in ("poll", "quiz"):
        first = r
        for oi, opt in enumerate(it["options"]):
            cell(res, r, 1, opt, wrap=True)
            for j, c in enumerate(confs):
                n = sum(1 for a in c["items"][qi]["answers"] if oi in a["value"])
                cell(res, r, 2 + j, n, fmt="0", align="center")
            tot = f"{get_column_letter(2 + N)}{r}"
            cell(res, r, 2 + N, f"=SUM({get_column_letter(2)}{r}:{get_column_letter(1 + N)}{r})", fmt="0", bold=True, align="center")
            cell(res, r, 3 + N, f"=IFERROR({tot}/{get_column_letter(2 + N)}${resp_row},0)", fmt="0.0%", align="center")
            r += 1
        if [o.strip() for o in it["options"]] == [str(k) for k in range(1, 11)]:
            energy_rows[qi] = (first, r - 1)
    elif t == "wordcloud":
        counts = [Counter(word(w) for a in c["items"][qi]["answers"] for w in a["value"] if word(w)) for c in confs]
        total = Counter()
        for cnt in counts:
            total.update(cnt)
        top = [w for w, _ in total.most_common(15)]
        for w in top:
            cell(res, r, 1, w)
            for j, cnt in enumerate(counts):
                cell(res, r, 2 + j, cnt.get(w, 0), fmt="0", align="center")
            cell(res, r, 2 + N, f"=SUM({get_column_letter(2)}{r}:{get_column_letter(1 + N)}{r})", fmt="0", bold=True, align="center")
            cell(res, r, 3 + N, f"=IFERROR({get_column_letter(2 + N)}{r}/{get_column_letter(2 + N)}${resp_row},0)", fmt="0.0%", align="center")
            r += 1
        cell(res, r, 1, f"… {max(0, len(total) - len(top))} autres mots : voir l’onglet « Mots »", fill=None)
        res.cell(row=r, column=1).font = F(italic=True, color="64748B")
        r += 1
    r += 1

# ---------------------------------------------------------------- Mots
mots = wb.create_sheet("Mots")
mots["A1"] = "Nuages de mots — tous les mots, par question"
mots["A1"].font = F(bold=True, size=14)
head(mots, 3, ["Question", "Mot"] + names + ["Total"], [46, 30] + [12] * N + [10])
mots.freeze_panes = "C4"
r = 4
for qi, it in enumerate(items):
    if it["type"] != "wordcloud":
        continue
    counts = [Counter(word(w) for a in c["items"][qi]["answers"] for w in a["value"] if word(w)) for c in confs]
    total = Counter()
    for cnt in counts:
        total.update(cnt)
    for w, _ in sorted(total.items(), key=lambda x: (-x[1], x[0])):
        cell(mots, r, 1, f"Q{qi + 1} · {it['title']}")
        cell(mots, r, 2, w)
        for j, cnt in enumerate(counts):
            cell(mots, r, 3 + j, cnt.get(w, 0), fmt="0", align="center")
        cell(mots, r, 3 + N, f"=SUM(C{r}:{get_column_letter(2 + N)}{r})", fmt="0", bold=True, align="center")
        r += 1

# ---------------------------------------------------------------- Participants
part = wb.create_sheet("Participants")
part["A1"] = "Participants — prénoms tels que saisis sur les téléphones"
part["A1"].font = F(bold=True, size=14)
head(part, 3, ["Animateur", "Code", "Prénom", "Arrivée", "Questions répondues"], [14, 12, 34, 18, 20])
part.freeze_panes = "A4"
r = 4
part_first = r
for c in confs:
    per = Counter(a["pid"] for it in c["items"] for a in it["answers"])
    for p in sorted(c["people"], key=lambda p: p["joinedAt"]):
        cell(part, r, 1, c["who"])
        cell(part, r, 2, c["label"], align="center")
        cell(part, r, 3, p["name"])
        cell(part, r, 4, when(p["joinedAt"]), fmt="dd/mm/yyyy hh:mm", align="center")
        cell(part, r, 5, per.get(p["id"], 0), fmt="0", align="center")
        r += 1
part_last = r - 1
part.auto_filter.ref = f"A3:E{part_last}"

# ---------------------------------------------------------------- Synthèse (en premier)
syn = wb.create_sheet("Synthèse", 0)
syn["A1"] = "Recharge tes 3 batteries · ICF & Edvantis — synthèse des conférences"
syn["A1"].font = F(bold=True, size=14)
syn["A2"] = f"Données extraites du site Live le {datetime.now(CASA):%d/%m/%Y à %H:%M} (heure du Maroc)."
syn["A2"].font = F(italic=True, color="64748B")
hdr = ["Animateur", "Code", "Date", "Inscrits", "Prénoms distincts", "Réponses", "Questions jouées",
       "Réponses par question", "Énergie avant pause (/10)", "Énergie après pause (/10)", "Évolution"]
head(syn, 4, hdr, [14, 11, 18, 11, 13, 12, 12, 14, 15, 15, 12])
syn.row_dimensions[4].height = 32
e_keys = sorted(energy_rows)  # avant puis après la pause active
r = 5
for j, c in enumerate(confs):
    col = get_column_letter(2 + j)  # colonne de la conférence dans « Résultats »
    first_join = min(p["joinedAt"] for p in c["people"]) if c["people"] else c["createdAt"]
    played = sum(1 for it in c["items"] if it["answers"])
    cell(syn, r, 1, c["who"], bold=True)
    cell(syn, r, 2, c["label"], align="center")
    cell(syn, r, 3, when(first_join), fmt="dd/mm/yyyy hh:mm", align="center")
    cell(syn, r, 4, f'=COUNTIF(Participants!$A${part_first}:$A${part_last},A{r})', fmt="0", align="center")
    cell(syn, r, 5, len({norm(p["name"]) for p in c["people"]}), fmt="0", align="center")
    resp = "+".join(f"'Résultats'!{col}{answered_cells[qi]}" for qi in answered_cells)
    cell(syn, r, 6, f"={resp}", fmt="0", align="center")
    cell(syn, r, 7, played, fmt="0", align="center")
    cell(syn, r, 8, f"=IFERROR(F{r}/G{r},0)", fmt="0.0", align="center")
    for k, qi in enumerate(e_keys[:2]):
        a, b = energy_rows[qi]
        # moyenne pondérée : choix « 1 » à « 10 » × nombre de réponses
        mean = f"=IFERROR(SUMPRODUCT(ROW('Résultats'!{col}{a}:{col}{b})-ROW('Résultats'!{col}{a})+1,'Résultats'!{col}{a}:{col}{b})/SUM('Résultats'!{col}{a}:{col}{b}),\"-\")"
        cell(syn, r, 9 + k, mean, fmt="0.0", align="center")
    cell(syn, r, 11, f'=IFERROR(J{r}-I{r},"-")', fmt="+0.0;-0.0;0.0", align="center")
    r += 1
last = r - 1
cell(syn, r, 1, "Total", bold=True, fill=TOTAL_FILL)
cell(syn, r, 2, None, fill=TOTAL_FILL)
cell(syn, r, 3, None, fill=TOTAL_FILL)
for col in "DEF":
    cell(syn, r, "ABCDEFGHIJK".index(col) + 1, f"=SUM({col}5:{col}{last})", fmt="0", bold=True, fill=TOTAL_FILL, align="center")
cell(syn, r, 7, f"=MAX(G5:G{last})", fmt="0", bold=True, fill=TOTAL_FILL, align="center")
cell(syn, r, 8, f"=IFERROR(AVERAGE(H5:H{last}),0)", fmt="0.0", bold=True, fill=TOTAL_FILL, align="center")
for k, colL in enumerate("IJ"):
    cell(syn, r, 9 + k, f'=IFERROR(AVERAGE({colL}5:{colL}{last}),"-")', fmt="0.0", bold=True, fill=TOTAL_FILL, align="center")
cell(syn, r, 11, f'=IFERROR(J{r}-I{r},"-")', fmt="+0.0;-0.0;0.0", bold=True, fill=TOTAL_FILL, align="center")

notes = [
    "Inscrits : téléphones qui ont rejoint la salle. Une personne qui se reconnecte depuis un autre navigateur compte deux fois.",
    "Prénoms distincts : estimation des personnes réelles (même prénom, sans tenir compte des majuscules et accents). Inclut les essais de Dan et des coachs.",
    "Questions jouées : questions qui ont reçu au moins une réponse (17 au programme).",
    "Énergie avant / après : moyenne des notes de 1 à 10 données juste avant et juste après la pause active (questions de l’onglet « Résultats »).",
    "Totaux : moyennes simples des conférences pour l’énergie.",
]
r += 2
syn.cell(row=r, column=1, value="Comment lire ce tableau").font = F(bold=True)
for n in notes:
    r += 1
    c = syn.cell(row=r, column=1, value="• " + n)
    c.font = F(color="475569")
syn.freeze_panes = "B5"

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("ok", OUT)
